import openpyxl

wb = openpyxl.load_workbook('prereq_data/VRCF Student DB.xlsx')
sheet = wb.active
headers = [sheet.cell(row=1, column=c).value for c in range(1, 18)]
for r in range(2, sheet.max_row + 1):
    for c in range(1, 18):
        val = str(sheet.cell(row=r, column=c).value or '')
        if len(val) > 25 and c in (3, 6, 7, 10, 12, 13, 14):
            print(f"Row {r}, Col {c} ({headers[c-1]}): len={len(val)} '{val}'")
